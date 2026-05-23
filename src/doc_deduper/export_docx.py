from __future__ import annotations

import csv
import textwrap
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .models import DocumentSegment
from .text import readable_text


def sort_key(segment: DocumentSegment) -> tuple[int, int, int, str]:
    decision = segment.decision_date or date(segment.year or 1, 1, 1)
    number = int(segment.number) if segment.number and segment.number.isdigit() else 0
    return (decision.year, decision.toordinal(), number, segment.title)


class Linker:
    def __init__(self) -> None:
        self.bookmark_id = 1

    def bookmark(self, paragraph, name: str) -> None:
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self.bookmark_id))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self.bookmark_id))
        self.bookmark_id += 1
        paragraph._p.insert(0, start)
        paragraph._p.append(end)


def add_internal_link(paragraph, text: str, anchor: str, *, bold: bool = False) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    if bold:
        rpr.append(OxmlElement("w:b"))
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_border(paragraph, color: str = "B7C9D8", size: str = "8") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = pbdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        pbdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)


def set_shading(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(4)

    for style_name, size, before, after, color in [
        ("Title", 22, 0, 12, "1F2933"),
        ("Heading 1", 14, 10, 7, "102A43"),
        ("Heading 2", 11, 8, 4, "334E68"),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name, size, color in [
        ("Deduper Meta", 8.5, "334E68"),
        ("Deduper Index", 9, "1F2933"),
        ("Deduper Divider", 8.5, "102A43"),
    ]:
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = styles["Normal"]
            style.font.name = "Arial"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string(color)
            style.paragraph_format.space_after = Pt(3)
    styles["Deduper Index"].paragraph_format.left_indent = Cm(0.35)
    styles["Deduper Index"].paragraph_format.first_line_indent = Cm(-0.35)


def split_line(line: str, max_chars: int) -> list[str]:
    line = readable_text(line)
    if len(line) <= max_chars:
        return [line]
    return textwrap.wrap(line, width=max_chars, break_long_words=False, break_on_hyphens=False) or [line]


def write_manifest(path: Path, segments: list[DocumentSegment]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["ordinal", "title", "source_kept", "number", "year", "all_sources"])
        for ordinal, segment in enumerate(segments, 1):
            writer.writerow(
                [
                    ordinal,
                    readable_text(segment.title),
                    segment.source_ref,
                    segment.number or "",
                    segment.year or "",
                    "; ".join(segment.duplicate_sources),
                ]
            )


def export_word(
    path: Path,
    segments: list[DocumentSegment],
    *,
    duplicates_removed: int,
    max_paragraph_chars: int = 2800,
) -> None:
    doc = Document()
    style_document(doc)
    linker = Linker()

    title = doc.add_paragraph(style="Title")
    title.add_run("Merged documents without duplicates")
    subtitle = doc.add_paragraph(
        f"{len(segments)} unique records | {duplicates_removed} duplicates removed",
        style="Deduper Meta",
    )
    set_shading(subtitle, "EEF5FB")

    summary = doc.add_heading("Clickable summary", level=1)
    linker.bookmark(summary, "summary")
    intro = doc.add_paragraph("Click an item to jump to it. Each record has a return link.", style="Deduper Meta")
    set_border(intro, "D9E2EC", "6")

    current_year: int | None = None
    for ordinal, segment in enumerate(segments, 1):
        if segment.year != current_year:
            current_year = segment.year
            year = doc.add_paragraph(str(current_year or "No year"), style="Deduper Divider")
            set_border(year, "BCCCDC", "6")
        row = doc.add_paragraph(style="Deduper Index")
        add_internal_link(row, f"{ordinal:03d}. {readable_text(segment.title)}", f"doc_{ordinal:03d}")
        meta = row.add_run(f" | source: {segment.source_ref}")
        meta.font.size = Pt(8)
        meta.font.color.rgb = RGBColor.from_string("627D98")

    for ordinal, segment in enumerate(segments, 1):
        doc.add_page_break()
        heading = doc.add_heading(f"{ordinal}. {readable_text(segment.title)}", level=1)
        linker.bookmark(heading, f"doc_{ordinal:03d}")
        set_border(heading, "829AB1", "10")

        meta = doc.add_paragraph(style="Deduper Meta")
        meta.add_run(f"Kept source: {segment.source_ref}").bold = True
        if len(segment.duplicate_sources) > 1:
            meta.add_run(" | Duplicates: " + "; ".join(segment.duplicate_sources))
        set_shading(meta, "F3F6F8")

        back = doc.add_paragraph(style="Deduper Meta")
        add_internal_link(back, "Back to summary", "summary")

        for raw_line in segment.lines:
            line = readable_text(raw_line)
            if not line or line == readable_text(segment.title):
                continue
            if line.upper() in {
                "INTESTAZIONE",
                "SVOLGIMENTO DEL PROCESSO",
                "MOTIVI DELLA DECISIONE",
                "RITENUTO IN FATTO",
                "CONSIDERATO IN DIRITTO",
                "P.Q.M.",
                "MASSIMA",
            }:
                doc.add_heading(line, level=2)
                continue
            for piece in split_line(line, max_paragraph_chars):
                paragraph = doc.add_paragraph(piece)
                if piece.upper() in {
                    "REPUBBLICA ITALIANA",
                    "IN NOME DEL POPOLO ITALIANO",
                    "LA CORTE SUPREMA DI CASSAZIONE",
                    "SENTENZA",
                }:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.bold = True

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def export_google_docs_chunks(
    out_dir: Path,
    segments: list[DocumentSegment],
    *,
    duplicates_removed: int,
    char_limit: int = 760000,
) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    parts: list[list[tuple[int, DocumentSegment, str]]] = []
    current: list[tuple[int, DocumentSegment, str]] = []
    current_chars = 0
    for ordinal, segment in enumerate(segments, 1):
        text = segment_as_text(ordinal, segment)
        if current and current_chars + len(text) > char_limit:
            parts.append(current)
            current = []
            current_chars = 0
        current.append((ordinal, segment, text))
        current_chars += len(text)
    if current:
        parts.append(current)

    paths: list[Path] = []
    total = len(parts)
    for part_no, part in enumerate(parts, 1):
        first, last = part[0][0], part[-1][0]
        doc = Document()
        style_document(doc)
        doc.add_paragraph(
            f"Merged documents - part {part_no:02d} of {total:02d}",
            style="Title",
        )
        doc.add_paragraph(
            f"Records {first}-{last} | {len(segments)} unique | {duplicates_removed} duplicates removed",
            style="Deduper Meta",
        )
        for ordinal, segment, _ in part:
            doc.add_page_break()
            doc.add_heading(f"{ordinal}. {readable_text(segment.title)}", level=1)
            doc.add_paragraph(f"Kept source: {segment.source_ref}", style="Deduper Meta")
            for line in segment.lines:
                line = readable_text(line)
                if line and line != readable_text(segment.title):
                    for piece in split_line(line, 2500):
                        doc.add_paragraph(piece)
        path = out_dir / f"merged-part-{part_no:02d}-of-{total:02d}.docx"
        doc.save(path)
        paths.append(path)
    return paths


def segment_as_text(ordinal: int, segment: DocumentSegment) -> str:
    lines = [
        "=" * 72,
        f"{ordinal}. {readable_text(segment.title)}",
        f"Kept source: {segment.source_ref}",
        "=" * 72,
        "",
    ]
    lines.extend(readable_text(line) for line in segment.lines if readable_text(line))
    return "\n".join(lines) + "\n"
